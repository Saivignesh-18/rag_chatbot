"""Document parsing: extract text from PDF, DOCX and TXT.

For PDFs, text is returned per page so that page numbers can be preserved for
source citations. DOCX/TXT have no reliable page structure, so their text is
returned as a single segment with ``page_number = None``.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from app.core.exceptions import CorruptedFileError, NoTextExtractedError
from app.core.logging import get_logger
from app.utils.text_cleaning import clean_text

logger = get_logger(__name__)


@dataclass
class PageSegment:
    """A contiguous piece of text with an optional source page number."""

    text: str
    page_number: int | None = None


@dataclass
class ParsedDocument:
    segments: list[PageSegment] = field(default_factory=list)
    page_count: int | None = None

    @property
    def full_text(self) -> str:
        return "\n\n".join(seg.text for seg in self.segments if seg.text)

    @property
    def total_characters(self) -> int:
        return sum(len(seg.text) for seg in self.segments)


def parse_document(data: bytes, file_type: str) -> ParsedDocument:
    """Parse raw bytes into cleaned page segments based on ``file_type``."""
    if file_type == "pdf":
        parsed = _parse_pdf(data)
    elif file_type == "docx":
        parsed = _parse_docx(data)
    elif file_type == "txt":
        parsed = _parse_txt(data)
    else:  # pragma: no cover - guarded earlier by validation
        raise CorruptedFileError(f"Unsupported file type: {file_type}")

    # Drop empty segments after cleaning.
    parsed.segments = [s for s in parsed.segments if s.text and s.text.strip()]

    if not parsed.segments or parsed.total_characters == 0:
        raise NoTextExtractedError()

    return parsed


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, OSError, ValueError) as exc:
        raise CorruptedFileError("The PDF could not be opened. It may be corrupted.") from exc

    if getattr(reader, "is_encrypted", False):
        # Try empty-password decryption; fail clearly if that does not work.
        try:
            if reader.decrypt("") == 0:  # 0 => failed
                raise CorruptedFileError("The PDF is password-protected.")
        except Exception as exc:  # noqa: BLE001
            raise CorruptedFileError("The PDF is password-protected.") from exc

    segments: list[PageSegment] = []
    for index, page in enumerate(reader.pages):
        try:
            raw = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - a single bad page shouldn't kill the doc
            logger.warning("Failed to extract text from PDF page %d: %s", index + 1, type(exc).__name__)
            raw = ""
        cleaned = clean_text(raw)
        segments.append(PageSegment(text=cleaned, page_number=index + 1))

    return ParsedDocument(segments=segments, page_count=len(reader.pages))


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises various errors
        raise CorruptedFileError("The DOCX file could not be opened. It may be corrupted.") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]

    # Include table cell text, which is common in requirements documents.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    cleaned = clean_text("\n".join(parts))
    return ParsedDocument(segments=[PageSegment(text=cleaned, page_number=None)], page_count=None)


def _parse_txt(data: bytes) -> ParsedDocument:
    text: str | None = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")

    cleaned = clean_text(text)
    return ParsedDocument(segments=[PageSegment(text=cleaned, page_number=None)], page_count=None)
